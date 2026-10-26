"""Public TikTok page counts pushed by the TK 看板 browser extension.

A separate SQLite file next to the sales database. Public page counts and the
TikTok Shop exports measure different things: on ten fully covered videos
checked 2026-10-17, Shop "Video views" ran at 0.4–0.76x the public play count,
while likes, comments and shares nearly matched. So snapshots are kept apart
and shown beside the Shop figures, never summed into them or used for the
manual review classification.

Each push keeps one row per video and observation; a push repeating the latest
counts only moves that row's last-seen time.
"""
from __future__ import annotations

import re
import secrets
import sqlite3
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path

COUNTS = ("views", "likes", "comments", "shares", "saves")
MAX_RECORDS = 2000
MAX_COUNT = 10 ** 13
SCHEMA = f"""
CREATE TABLE IF NOT EXISTS snapshots (
    video_id TEXT, observed_at TEXT, {", ".join(f"{c} INTEGER" for c in COUNTS)},
    author TEXT, title TEXT, published_at TEXT, last_seen_at TEXT, received_at TEXT,
    PRIMARY KEY (video_id, observed_at)
);
CREATE TABLE IF NOT EXISTS meta (name TEXT PRIMARY KEY, value TEXT);
"""


class PublicStatsError(ValueError):
    pass


def _iso(moment):
    return moment.astimezone(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def _count(value):
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value != value or not 0 <= value <= MAX_COUNT:
        raise PublicStatsError("数值无效")
    return int(round(value))


def _observed(value, now):
    """The extension's observation time (ms since epoch); a missing or implausible one falls back to now."""
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value == value:
        try:
            moment = datetime.fromtimestamp(value / 1000, timezone.utc)
        except (OverflowError, OSError, ValueError):
            return now
        if now - timedelta(days=30) <= moment <= now + timedelta(minutes=5):
            return moment
    return now


def clean_record(raw, now):
    """One extension video record reduced to the counted fields; raises PublicStatsError when unusable."""
    if not isinstance(raw, dict):
        raise PublicStatsError("记录格式无效")
    video_id = str(raw.get("id") or "")
    if not re.fullmatch(r"\d{6,25}", video_id):
        raise PublicStatsError("视频 ID 无效")
    counts = {name: _count(raw.get(name)) for name in COUNTS}
    if all(value is None for value in counts.values()):
        raise PublicStatsError("没有任何数值")
    author = str(raw.get("author") or "").lstrip("@")
    title = raw.get("title") if isinstance(raw.get("title"), str) else ""
    published = raw.get("publishedAt") if isinstance(raw.get("publishedAt"), str) else ""
    return {
        "video_id": video_id, "observed_at": _iso(_observed(raw.get("observedAt"), now)), **counts,
        "author": author if re.fullmatch(r"[\w.\-]{1,100}", author) else "",
        "title": re.sub(r"[\x00-\x1f\x7f]+", " ", title).strip()[:300],
        "published_at": published[:40] if re.fullmatch(r"\d{4}-\d{2}-\d{2}T[\d:.]+Z?", published[:40]) else "",
    }


class PublicStatsStore:
    def __init__(self, runtime_root):
        self.database = Path(runtime_root) / "performance" / "public-stats.sqlite"
        self.lock = threading.Lock()

    def connect(self):
        self.database.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(self.database, timeout=10)
        connection.row_factory = sqlite3.Row
        connection.executescript(SCHEMA)
        return connection

    # ---- pairing key --------------------------------------------------------
    def link_key(self):
        with self.lock:
            connection = self.connect()
            try:
                row = connection.execute("SELECT value FROM meta WHERE name = 'link_key'").fetchone()
                if row:
                    return row[0]
                key = secrets.token_urlsafe(24)
                with connection:
                    connection.execute("INSERT INTO meta VALUES ('link_key', ?)", (key,))
                return key
            finally:
                connection.close()

    def rotate_key(self):
        key = secrets.token_urlsafe(24)
        with self.lock:
            connection = self.connect()
            try:
                with connection:
                    connection.execute("INSERT OR REPLACE INTO meta VALUES ('link_key', ?)", (key,))
            finally:
                connection.close()
        return key

    def check_key(self, offered):
        return bool(offered) and secrets.compare_digest(str(offered), self.link_key())

    def signature(self):
        """Read-only change marker, including committed writes still in SQLite's WAL."""
        values = []
        for path in (self.database, Path(str(self.database) + "-wal")):
            try:
                stat = path.stat()
                values.append((str(path), stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns, getattr(stat, "st_ino", None)))
            except OSError:
                values.append((str(path), None))
        return tuple(values)

    # ---- snapshots ----------------------------------------------------------
    def record(self, videos, is_known, now=None):
        """Save the pushed records whose video the workbench knows (is_known: ids -> known subset).
        Returns how many were saved, unchanged, ignored as unknown, and rejected as invalid."""
        if not isinstance(videos, list) or not videos:
            raise PublicStatsError("没有可推送的视频数据。")
        if len(videos) > MAX_RECORDS:
            raise PublicStatsError(f"一次最多推送 {MAX_RECORDS} 条视频。")
        now = now or datetime.now(timezone.utc)
        cleaned, invalid = {}, 0
        for raw in videos:
            try:
                record = clean_record(raw, now)
            except PublicStatsError:
                invalid += 1
                continue
            previous = cleaned.get(record["video_id"])
            if not previous or record["observed_at"] >= previous["observed_at"]:
                cleaned[record["video_id"]] = record
        known = set(is_known(list(cleaned))) if cleaned else set()
        ignored = sorted(set(cleaned) - known)
        saved = unchanged = 0
        received = _iso(now)
        with self.lock:
            connection = self.connect()
            try:
                with connection:
                    for video_id in sorted(known):
                        record = cleaned[video_id]
                        latest = connection.execute(
                            "SELECT * FROM snapshots WHERE video_id = ? ORDER BY observed_at DESC LIMIT 1", (video_id,)).fetchone()
                        if latest and max(latest["observed_at"], latest["last_seen_at"] or "") > record["observed_at"]:
                            unchanged += 1  # an older page read than the one already kept
                            continue
                        if latest and all(latest[name] == record[name] for name in COUNTS):
                            connection.execute("UPDATE snapshots SET last_seen_at = ? WHERE video_id = ? AND observed_at = ?",
                                               (record["observed_at"], video_id, latest["observed_at"]))
                            unchanged += 1
                            continue
                        connection.execute(
                            f"INSERT OR REPLACE INTO snapshots VALUES ({','.join('?' * (len(COUNTS) + 7))})",
                            (video_id, record["observed_at"], *(record[name] for name in COUNTS), record["author"],
                             record["title"], record["published_at"], record["observed_at"], received))
                        saved += 1
            finally:
                connection.close()
        return {"saved": saved, "unchanged": unchanged, "ignored": len(ignored), "ignored_ids": ignored[:20], "invalid": invalid}

    def latest(self, video_ids):
        """Video ID -> its latest public snapshot, for just these videos."""
        ids = sorted({str(v) for v in video_ids if v})
        if not ids or not self.database.is_file():
            return {}
        result = {}
        with self.lock:
            connection = self.connect()
            try:
                for start in range(0, len(ids), 500):
                    chunk = ids[start:start + 500]
                    for row in connection.execute(f"""
                            SELECT s.*, (SELECT COUNT(*) FROM snapshots c WHERE c.video_id = s.video_id) snapshot_count
                            FROM snapshots s WHERE s.video_id IN ({','.join('?' * len(chunk))})
                              AND s.observed_at = (SELECT MAX(observed_at) FROM snapshots m WHERE m.video_id = s.video_id)""", chunk):
                        item = dict(row)
                        result[item.pop("video_id")] = item
            finally:
                connection.close()
        return result


EXPORT_COLUMNS = (
    ("视频 ID", 22), ("账号", 18), ("标题", 46), ("产品", 10), ("发布日期", 12), ("参考价值", 18), ("首两周播放", 12),
    ("后台累计播放", 13), ("公开页播放", 13), ("公开页抓取时间", 18), ("公开页点赞", 11), ("公开页评论", 11),
    ("公开页转发", 11), ("公开页收藏", 11), ("订单", 8), ("售出件数", 9), ("归因 GMV ($)", 12), ("商品点击", 10),
    ("点击率", 9), ("CTOR", 9), ("后台数据覆盖", 24), ("视频链接", 30),
)
EXPORT_NOTES = (
    "后台累计播放：销售数据表（TikTok Shop 导出）已导入日表/周表覆盖内的播放加总；未覆盖的日期没有计入。",
    "公开页播放：TK 看板插件从 TikTok 公开页推送的累计播放，以「公开页抓取时间」为准；大数可能已被 TikTok 取整。空白表示还没推送过。",
    "参考价值来自人工复盘分类；未复盘显示待复盘，已复盘未分类显示待分类。首两周播放、点击率、CTOR、订单与 GMV 来自销售数据表；点击率 = 商品点击 ÷ 后台累计播放。",
    "本表由TK工作台生成，不是后台原始导出；不要放进销售数据文件夹，以免被当成导出表导入。",
)


def export_workbook(videos, public, target):
    """Write the chosen videos' export and public figures side by side to an .xlsx file."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "公开播放"
    sheet.append([name for name, _ in EXPORT_COLUMNS])
    for index, (_, width) in enumerate(EXPORT_COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="E8F1FC")
        cell.alignment = Alignment(vertical="center")
    for video in videos:
        snapshot = public.get(video["video_id"]) or {}
        seen = snapshot.get("last_seen_at") or snapshot.get("observed_at") or ""
        if seen:
            seen = datetime.fromisoformat(seen.replace("Z", "+00:00")).astimezone().strftime("%Y-%m-%d %H:%M")
        shop_views = video.get("views_shop", video.get("views")) if video.get("views_source") == "public" else video.get("views")
        coverage = f"{video['first_data']} 至 {video['last_data']}" if video.get("first_data") else ""
        sheet.append([
            video["video_id"], video.get("account") or "", video.get("headline") or video.get("title") or "", video.get("product") or "",
            video.get("posted_on") or "", (video.get("review_status") or {}).get("label") or "待复盘", video.get("window_views"), shop_views,
            snapshot.get("views"), seen, snapshot.get("likes"), snapshot.get("comments"), snapshot.get("shares"), snapshot.get("saves"),
            video.get("orders"), video.get("items"), video.get("gmv"), video.get("clicks"),
            None if video.get("ctr_pct") is None else video["ctr_pct"] / 100,
            None if video.get("ctor_pct") is None else video["ctor_pct"] / 100, coverage,
            f"https://www.tiktok.com/@{video.get('account') or '_'}/video/{video['video_id']}",
        ])
    for row in sheet.iter_rows(min_row=2):
        for cell in row[6:9] + row[10:16] + row[17:18]:
            cell.number_format = "#,##0"
        row[16].number_format = "#,##0.00"
        row[18].number_format = row[19].number_format = "0.00%"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    notes = book.create_sheet("说明")
    notes.column_dimensions["A"].width = 110
    for line in EXPORT_NOTES:
        notes.append([line])
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return target
