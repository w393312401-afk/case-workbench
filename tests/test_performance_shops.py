"""Shops: files are told apart by folder, nothing crosses between shops, gaps count per shop, photo exports import."""
from __future__ import annotations

import sqlite3
import tempfile
import unittest
from pathlib import Path

import test_performance as exports
import test_server as fixture
from performance_store import CREATOR_TEXT, METRICS, TOTAL_METRICS, VIDEO_TEXT, PerformanceStore, parse_export, read_export

LINKED = exports.LINKED
CROSS_VIDEO = "7600000000000000002"  # an ID the export fixture already knows
PHOTO_HEADER = ["Name of product", "Product ID", "Photo attributed GMV ($)", "Attributed SKU orders",
                "Product Impressions", "Click-through rate", "CTOR (SKU order)", "GPM ($)"]
PHOTO_ROWS = [
    ["SHARPAL 202H Precision Kitchen Knife Sharpening System", "1732157836537861037", "1094.32", "8", "18817", "0.0138", "0.0308", "58.16"],
    [None, "-1", "300.2", "6", "0", "0", "0", "0"],
    ["SHARPAL 6-Function Compact Pocket Knife", "1729474277203022765", "27.53", "2", "768", "0.0143", "0.1818", "35.85"],
]


def photo_export(path, start, end, rows=PHOTO_ROWS, shop_name="SHARPAL US"):
    import openpyxl
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([f"[Date Range]: {start} ~ {end}\n"])
    if shop_name:
        sheet.append([f"Shop Name: {shop_name}"])
    sheet.append([None])
    sheet.append(PHOTO_HEADER)
    for values in rows:
        sheet.append(values)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


ITEM_HEADER = ["Creator name", "Creator ID", "Photo info", "Photo item ID", "Time", "Products", "Photo attributed GMV ($)",
               "Photo attributed items sold", "Photo attributed items sold", "CTOR (SKU order)", "Views", "Product Impressions",
               "Click-through rate", "Likes", "Shares", "Comments", "GPM"]
# first "items sold" 3, second 9: only the first may be read
ITEM_ROWS = [
    ["sharpal_us", "1", "Kitchen knife tip #knife #tips", "7685397373245312269", "2026/09/14 06:27:38", "SHARPAL 198H Electric Sharpener(1729474345501823917)",
     "120.5", "3", "9", "4.00%", "643", "660", "2.18%", "5", "1", "2", "182.58"],
    ["bladesharptrevor", "2", None, "7688544785803660574", "2026/09/22 18:00:49", "SHARPAL 202H Precision Kit(1732157836537861037)",
     "0", "0", "0", "0.00%", "16", "15", "12.50%", "0", "0", "0", "0"],
]


def item_export(path, start, end, rows=ITEM_ROWS, shop_line="Shop: SHARPAL US\n"):
    import openpyxl
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([f"[Date Range]: {start} ~ {end}\n"])
    sheet.append([shop_line])
    sheet.append([None])
    sheet.append(ITEM_HEADER)
    for values in rows:
        sheet.append(values)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def totals(path, day, gmv):
    return exports.export(path, day, day, [[day, str(gmv), "90", "10", "5000", "20", "3", "2", "1", "3", "4000", "5100", "120", "0.0244", "0.0125"]],
                          exports.TOTALS_HEADER)


class ShopStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="performance-shops-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name) / "exports"
        local, cross = self.folder / "本土店", self.folder / "跨境店"
        exports.export(local / "视频" / "9月" / "第一周视频.xlsx", "2026-08-31", "2026-09-06", [exports.row(LINKED, 8000, 3, 200, 450.0)])
        exports.export(local / "视频" / "9月" / "第二周视频.xlsx", "2026-09-07", "2026-09-13", [exports.row(LINKED, 4000, 1, 100, 150.0)])
        totals(local / "销售" / "9月" / "第二周销售.xlsx", "2026-09-07", 100)
        # The cross-border shop opened later, exports the same week as the local one, and has its own video.
        exports.export(cross / "视频" / "第二周视频.xlsx", "2026-09-07", "2026-09-13", [exports.row(CROSS_VIDEO, 700, 2, 40, 80.0)])
        totals(cross / "销售" / "第二周销售.xlsx", "2026-09-07", 55)
        self.store = PerformanceStore(Path(self.temp.name) / "runtime", [self.folder])

    def test_folder_decides_the_shop_and_the_same_week_or_day_does_not_overwrite(self):
        result = self.store.sync()
        self.assertEqual(result["failed"], [])
        self.assertEqual(sorted((f["file"], f["shop"]) for f in result["imported"]),
                         [("第一周视频.xlsx", "local"), ("第二周视频.xlsx", "cross"), ("第二周视频.xlsx", "local"),
                          ("第二周销售.xlsx", "cross"), ("第二周销售.xlsx", "local")])
        connection = sqlite3.connect(self.store.database)
        try:
            # Both shops exported 9/7 (a totals row each) and the week 9/7–9/13 (a periods block each): all of it is kept.
            self.assertEqual(connection.execute("SELECT shop, gmv FROM totals_daily ORDER BY shop").fetchall(), [("cross", 55.0), ("local", 100.0)])
            self.assertEqual(connection.execute("SELECT shop, COUNT(*) FROM periods WHERE range_start = '2026-09-07' GROUP BY shop ORDER BY shop").fetchall(),
                             [("cross", 1), ("local", 1)])
        finally:
            connection.close()
        self.assertEqual([d["gmv"] for d in self.store.analytics("local")["shop_daily"]], [100.0])
        self.assertEqual([d["gmv"] for d in self.store.analytics("cross")["shop_daily"]], [55.0])
        self.assertEqual(sum(d["gmv"] for d in self.store.analytics()["shop_daily"]), 155.0)

    def test_each_shop_sees_only_its_own_videos_and_figures(self):
        self.store.sync()
        self.assertEqual([v["video_id"] for v in self.store.videos(shop="local")], [LINKED])
        self.assertEqual([v["video_id"] for v in self.store.videos(shop="cross")], [CROSS_VIDEO])
        self.assertEqual(len(self.store.videos()), 2)
        local = self.store.analytics("local")
        self.assertEqual([row[0] for row in local["videos"]], [LINKED])
        self.assertEqual(sum(r[7] for r in local["weekly"]), 600.0)
        self.assertEqual(sum(r[7] for r in self.store.analytics("cross")["weekly"]), 80.0)
        self.assertEqual(self.store.video(CROSS_VIDEO)["all_videos"], 1)  # ranked among its own shop's videos
        with self.assertRaises(ValueError):
            self.store.summary("nowhere")

    def test_gaps_and_coverage_are_worked_out_per_shop(self):
        self.store.sync()
        local, cross = self.store.summary("local"), self.store.summary("cross")
        self.assertEqual(local["coverage"], {"first": "2026-08-31", "last": "2026-09-13", "gaps": []})
        # The cross-border shop starts a week later: nothing before it is a gap, and its own weeks are complete.
        self.assertEqual(cross["coverage"], {"first": "2026-09-07", "last": "2026-09-13", "gaps": []})
        self.assertEqual([(w["week"], w["complete"]) for w in cross["weeks"]], [("2026-09-07", True)])
        self.assertEqual([(s["id"], s["files"]) for s in cross["shops"]], [("local", 3), ("cross", 2)])
        # Without a shop the weeks of both are merged, as before shops existed.
        self.assertEqual([w["week"] for w in self.store.summary()["weeks"]], ["2026-08-31", "2026-09-07"])

    def test_a_videos_status_is_judged_against_its_own_shops_last_day(self):
        # Local data runs to 9/13; a later cross-border export (to 10/4) must not turn the local video's open window into "closed".
        exports.export(self.folder / "本土店" / "视频" / "9月" / "第二周视频.xlsx", "2026-09-07", "2026-09-13",
                       [exports.row(LINKED, 4000, 1, 100, 150.0), exports.row("7600000000000000003", 900)])  # posted 9/9: window ends 9/20
        exports.export(self.folder / "跨境店" / "视频" / "第四周视频.xlsx", "2026-09-28", "2026-10-04", [exports.row(CROSS_VIDEO, 100)])
        self.store.sync()
        statuses = {v["video_id"]: v["status"] for v in self.store.videos()}
        self.assertEqual((statuses[LINKED], statuses["7600000000000000003"]), ("focus", "pending"))

    def test_a_file_in_the_wrong_type_folder_is_refused(self):
        misplaced = totals(self.folder / "本土店" / "视频" / "第三周销售.xlsx", "2026-09-14", 9)
        result = self.store.sync()
        failed = {f["file"]: f["error"] for f in result["failed"]}
        self.assertIn("销售表", failed[misplaced.name])
        self.assertEqual([d["date"] for d in self.store.analytics("local")["shop_daily"]], ["2026-09-07"])

    def test_a_forced_shop_overrides_the_folder(self):
        loose = totals(Path(self.temp.name) / "picked.xlsx", "2026-09-20", 7)
        result = self.store.import_paths([loose], shop="cross")
        self.assertEqual(result["imported"][0]["shop"], "cross")
        self.assertEqual([d["date"] for d in self.store.analytics("cross")["shop_daily"]], ["2026-09-20"])
        with self.assertRaises(Exception):
            self.store.import_paths([loose], shop="nowhere")
        self.assertEqual(self.store.import_paths([totals(Path(self.temp.name) / "plain.xlsx", "2026-09-21", 8)])["imported"][0]["shop"], "local")


class PhotoExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="performance-photo-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name) / "exports"
        self.store = PerformanceStore(Path(self.temp.name) / "runtime", [self.folder])

    def test_photo_table_is_read_past_the_shop_name_line(self):
        path = photo_export(self.folder / "本土店" / "图文" / "10月" / "第二周图文.xlsx", "2026-10-12", "2026-10-18")
        kind, start, end, records = read_export(path)
        self.assertEqual((kind, start, end, len(records)), ("photo", "2026-10-12", "2026-10-18", 3))
        self.assertAlmostEqual(records[0]["ctr_pct"], 1.38)
        self.assertEqual(parse_export(exports_rows(path))[4], "SHARPAL US")

    def test_import_and_read_per_shop(self):
        photo_export(self.folder / "本土店" / "图文" / "第一周图文.xlsx", "2026-10-05", "2026-10-11")
        photo_export(self.folder / "本土店" / "图文" / "第二周图文.xlsx", "2026-10-12", "2026-10-18")
        photo_export(self.folder / "跨境店" / "图文" / "第一周图文.xlsx", "2026-10-12", "2026-10-18",
                     [["SHARPAL 202H Precision Kitchen Knife Sharpening System", "1732157836537861037", "40", "1", "900", "0.02", "0.05", "44"]],
                     shop_name="SHARPAL GLOBAL")
        result = self.store.sync()
        self.assertEqual((result["failed"], [f["kind"] for f in result["imported"]]), ([], ["photo_periods"] * 3))
        data = self.store.photos("local")
        self.assertEqual([(p["start"], p["weekly"], p["shop_names"]) for p in data["periods"]],
                         [("2026-10-05", True, ["SHARPAL US"]), ("2026-10-12", True, ["SHARPAL US"])])
        names = {p[0]: p[1] for p in data["products"]}
        self.assertEqual(names, {"1732157836537861037": "202H", "-1": "未归属商品", "1729474277203022765": "6-Function Compact Pocket"})
        at = [p[0] for p in data["products"]].index("1732157836537861037")
        # 1,094.32 GMV, 8 orders, 18,817 impressions; clicks follow from the 1.38% rate.
        self.assertIn([at, 0, 1094.32, 8, 18817, round(18817 * 0.0138)], data["rows"])
        self.assertAlmostEqual(sum(r[2] for r in data["rows"]), 2 * (1094.32 + 300.2 + 27.53))
        cross = self.store.photos("cross")
        self.assertEqual((len(cross["periods"]), cross["periods"][0]["shop_names"], len(cross["rows"])), (1, ["SHARPAL GLOBAL"], 1))
        self.assertEqual(self.store.summary("cross")["shops"][1]["shop_names"], ["SHARPAL GLOBAL"])
        self.assertEqual(self.store.summary()["shops"][0]["photo_files"], 2)

    def test_reimport_replaces_the_range_and_a_month_never_adds_onto_its_weeks(self):
        week = photo_export(self.folder / "本土店" / "图文" / "第二周图文.xlsx", "2026-10-12", "2026-10-18")
        photo_export(self.folder / "本土店" / "图文" / "10月图文.xlsx", "2026-10-01", "2026-10-31")
        self.store.sync()
        data = self.store.photos("local")
        self.assertEqual([p["start"] for p in data["periods"]], ["2026-10-12"])
        self.assertEqual([e["start"] for e in data["excluded_periods"]], ["2026-10-01"])
        photo_export(week, "2026-10-12", "2026-10-18", PHOTO_ROWS[:1])
        self.store.sync()
        self.assertEqual(len(self.store.photos("local")["rows"]), 1)


def exports_rows(path):
    import openpyxl
    return list(openpyxl.load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True))


class PhotoItemTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="performance-items-")
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name) / "exports"
        self.store = PerformanceStore(Path(self.temp.name) / "runtime", [self.folder])

    def test_the_per_post_table_is_read_with_its_first_items_sold_column_and_the_shop_line(self):
        path = item_export(self.folder / "本土店" / "图文" / "9月" / "图文明细.xlsx", "2026-09-01", "2026-09-30")
        kind, start, end, records, shop_name = parse_export(exports_rows(path))
        self.assertEqual((kind, start, end, len(records), shop_name), ("photo_items", "2026-09-01", "2026-09-30", 2, "SHARPAL US"))
        first = records[0]
        self.assertEqual((first["item_id"], first["sold"], first["views"], first["posted_at"]), ("7685397373245312269", 3.0, 643.0, "2026-09-14 06:27:38"))
        self.assertAlmostEqual(first["ctr_pct"], 2.18)
        self.assertAlmostEqual(first["ctor_pct"], 4.0)

    def test_import_per_shop_clicks_are_worked_back_and_a_new_export_replaces_the_range(self):
        month = item_export(self.folder / "本土店" / "图文" / "9月" / "图文明细.xlsx", "2026-09-01", "2026-09-30")
        item_export(self.folder / "跨境店" / "图文" / "图文明细.xlsx", "2026-09-01", "2026-09-30", ITEM_ROWS[1:], shop_line="Shop Name: SHARPAL GLOBAL")
        result = self.store.sync()
        self.assertEqual((result["failed"], sorted((f["kind"], f["shop"]) for f in result["imported"])),
                         ([], [("photo_items", "cross"), ("photo_items", "local")]))
        data = self.store.photo_items("local")
        self.assertEqual([i[0] for i in data["items"]], ["7685397373245312269", "7688544785803660574"])
        self.assertEqual(data["items"][0][1:], ["sharpal_us", "Kitchen knife tip", "2026-09-14", ["198H"]])
        self.assertEqual(data["items"][1][2], "")  # a post with no text stays empty rather than inventing one
        # views 643, impressions 660, clicks 660 x 2.18% ~ 14, sold 3, GMV 120.5, likes 5
        self.assertEqual(data["rows"][0], [0, 0, 643, 660, round(660 * 0.0218), 3, 120.5, 5, 1, 2])
        self.assertEqual(data["periods"][0]["shop_names"], ["SHARPAL US"])
        self.assertEqual(len(self.store.photo_items("cross")["rows"]), 1)
        self.assertEqual(self.store.photo_items("cross")["periods"][0]["shop_names"], ["SHARPAL GLOBAL"])
        self.assertEqual(self.store.summary()["shops"][0]["photo_files"], 1)
        item_export(month, "2026-09-01", "2026-09-30", ITEM_ROWS[:1])
        self.store.sync()
        self.assertEqual(len(self.store.photo_items("local")["rows"]), 1)

    def test_a_longer_export_never_adds_onto_an_overlapping_one(self):
        item_export(self.folder / "本土店" / "图文" / "9月明细.xlsx", "2026-09-01", "2026-09-30")
        item_export(self.folder / "本土店" / "图文" / "9月上半月.xlsx", "2026-09-01", "2026-09-15")
        self.store.sync()
        data = self.store.photo_items("local")
        self.assertEqual([p["end"] for p in data["periods"]], ["2026-09-30"])
        self.assertEqual([e["end"] for e in data["excluded_periods"]], ["2026-09-15"])
        self.assertEqual(len(data["rows"]), 2)

    def test_the_file_must_sit_in_the_photo_folder(self):
        item_export(self.folder / "本土店" / "视频" / "图文明细.xlsx", "2026-09-01", "2026-09-30")
        result = self.store.sync()
        self.assertEqual([f["error"].startswith("文件放在「视频」文件夹里") for f in result["failed"]], [True])


class MigrationTests(unittest.TestCase):
    def test_a_database_from_before_shops_keeps_its_rows_as_the_local_shops(self):
        temp = tempfile.TemporaryDirectory(prefix="performance-migrate-")
        self.addCleanup(temp.cleanup)
        store = PerformanceStore(Path(temp.name) / "runtime", [])
        store.root.mkdir(parents=True)
        video, metrics = ", ".join(f"{c} TEXT" for c in VIDEO_TEXT), ", ".join(f"{m} REAL" for m in METRICS)
        creator = ", ".join(f"{c} TEXT" for c in CREATOR_TEXT)
        old = f"""
        CREATE TABLE daily (video_id TEXT, date TEXT, {video}, {metrics}, diagnosis TEXT, source_file TEXT, PRIMARY KEY (video_id, date));
        CREATE TABLE periods (video_id TEXT, range_start TEXT, range_end TEXT, span_days INTEGER, {video}, {metrics}, diagnosis TEXT,
            source_file TEXT, PRIMARY KEY (video_id, range_start, range_end));
        CREATE TABLE creator_periods (video_id TEXT, range_start TEXT, range_end TEXT, span_days INTEGER, {creator}, {metrics},
            diagnosis TEXT, source_file TEXT, PRIMARY KEY (video_id, range_start, range_end));
        CREATE TABLE totals_daily (date TEXT PRIMARY KEY, {", ".join(f"{m} REAL" for m in TOTAL_METRICS)}, source_file TEXT);
        CREATE TABLE imports (source_file TEXT PRIMARY KEY, kind TEXT, range_start TEXT, range_end TEXT, rows INTEGER, modified TEXT, imported_at TEXT);
        """
        connection = sqlite3.connect(store.database)
        connection.executescript(old)
        self.assertNotIn("shop", [r[1] for r in connection.execute("PRAGMA table_info(daily)")])
        self.assertNotIn("shop", [r[1] for r in connection.execute("PRAGMA table_info(totals_daily)")])
        connection.execute("INSERT INTO totals_daily (date, gmv, source_file) VALUES ('2026-09-07', 12.5, 'old.xlsx')")
        connection.execute("INSERT INTO imports (source_file, kind) VALUES ('old.xlsx', 'totals_daily')")
        connection.commit()
        connection.close()
        migrated = store.connect()
        try:
            self.assertEqual([tuple(r) for r in migrated.execute("SELECT shop, date, gmv FROM totals_daily")], [("local", "2026-09-07", 12.5)])
            self.assertEqual([tuple(r) for r in migrated.execute("SELECT shop FROM imports")], [("local",)])
            # The key is now (shop, date): the same day for another shop fits.
            migrated.execute("INSERT INTO totals_daily (shop, date, gmv) VALUES ('cross', '2026-09-07', 1)")
            for table in ("daily", "periods", "creator_periods"):
                self.assertIn("shop", [r[1] for r in migrated.execute(f"PRAGMA table_info({table})")])
        finally:
            migrated.close()
        store.connect().close()  # running it again is harmless


class ShopHTTPTests(unittest.TestCase):
    setUp_ = fixture.WorkbenchHTTPTests.setUp
    stop_server = fixture.WorkbenchHTTPTests.stop_server
    request = fixture.WorkbenchHTTPTests.request

    def setUp(self):
        self.setUp_()
        self.folder = self.root / "exports"
        exports.export(self.folder / "本土店" / "视频" / "第一周视频.xlsx", "2026-08-31", "2026-09-06", [exports.row(LINKED, 8000, 3, 200, 450.0)])
        exports.export(self.folder / "跨境店" / "视频" / "第一周视频.xlsx", "2026-08-31", "2026-09-06", [exports.row(CROSS_VIDEO, 700, 2, 40, 80.0)])
        photo_export(self.folder / "本土店" / "图文" / "第二周图文.xlsx", "2026-10-12", "2026-10-18")
        self.server.app.performance.sources = [self.folder]

    def test_shop_parameter_scopes_every_performance_read_and_defaults_to_local(self):
        status, _, result = self.request("POST", "/api/performance/sync", {})
        self.assertEqual((status, result["failed"]), (200, []), result)
        ids = lambda shop: [v["video_id"] for v in self.request("GET", "/api/performance/videos" + shop)[2]["videos"]]
        self.assertEqual((ids(""), ids("?shop=local"), ids("?shop=cross")), ([LINKED], [LINKED], [CROSS_VIDEO]))
        self.assertEqual(self.request("GET", "/api/performance/summary?shop=cross")[2]["shop"], "cross")
        self.assertEqual([len(self.request("GET", f"/api/performance/analytics?shop={s}")[2]["videos"]) for s in ("local", "cross")], [1, 1])
        self.assertEqual(self.request("GET", "/api/performance/creators?shop=cross")[0], 200)
        item_export(self.folder / "本土店" / "图文" / "图文明细.xlsx", "2026-09-01", "2026-09-30")
        self.assertEqual(self.request("POST", "/api/performance/sync", {})[2]["failed"], [])
        self.server.app.production.public = lambda product=None: {"carousels": [
            {"id": "CAR-0001", "publication": {"post_id": "7688544785803660574"}, "publication_history": [{"post_id": "7685397373245312269"}]},
            {"id": "CAR-0002", "publication": None}]}
        items = self.request("GET", "/api/performance/photo-items")[2]
        self.assertEqual((len(items["rows"]), items["productions"]), (2, {"7688544785803660574": "CAR-0001", "7685397373245312269": "CAR-0001"}))
        self.assertEqual(self.request("GET", "/api/performance/photo-items?shop=cross")[2]["rows"], [])
        self.assertEqual(self.request("GET", "/api/performance/photo-items?shop=x")[0], 400)
        photos = self.request("GET", "/api/performance/photos")[2]
        self.assertEqual((photos["shop"], len(photos["rows"])), ("local", 3))
        self.assertEqual(self.request("GET", "/api/performance/photos?shop=cross")[2]["rows"], [])
        for bad in ("?shop=x", "?shop=local&shop=cross"):
            self.assertEqual(self.request("GET", "/api/performance/summary" + bad)[0], 400, bad)
        self.assertEqual(self.request("POST", "/api/performance/import", {"paths": ["x"], "shop": "x"})[0], 400)


if __name__ == "__main__":
    unittest.main()
